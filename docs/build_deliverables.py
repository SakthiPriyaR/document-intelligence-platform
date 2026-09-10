"""Build the submission presentation and PDF companion from repository content.

The script is intentionally dependency-light and is run with the presentation
generation dependencies installed in a temporary environment. It does not
read or write secrets and it does not touch application code.
"""

from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

INK = "102C32"
TEAL = "0F766E"
MINT = "EAF7F3"
PAPER = "F6F4EE"
MUTED = "66767A"
GOLD = "C48A38"
LINE = "D6E1DE"
WHITE = "FFFFFF"
RED = "A84232"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_fill(shape, value: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(value)


def set_line(shape, value: str, width: float = 1) -> None:
    shape.line.color.rgb = rgb(value)
    shape.line.width = Pt(width)


def add_text(slide, text: str, x, y, w, h, size=16, color=INK,
             bold=False, font="Aptos", align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_link(slide, label: str, url: str, x, y, w, h, size=12):
    """Add a visible, clickable hyperlink to the PowerPoint deck."""
    box = add_text(slide, label, x, y, w, h, size, TEAL, True)
    box.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=15, color=INK,
                   line_spacing=1.15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(size * (line_spacing - 1))
    return box


def add_bullets(slide, items, x, y, w, h, size=15, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.level = 0
        p.space_after = Pt(8)
    return box


def add_title(slide, title: str, kicker: str, number: int) -> None:
    add_text(slide, kicker.upper(), 0.65, 0.38, 5.5, 0.25, 9, TEAL, True)
    add_text(slide, title, 0.65, 0.70, 11.6, 0.52, 27, INK, True)
    add_text(slide, f"{number:02d}  |  DOCUMENT INTELLIGENCE PLATFORM", 9.45, 7.03,
             3.2, 0.2, 8, MUTED, True, align=PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.43),
                                  Inches(12.0), Inches(0.015))
    set_fill(rule, LINE)
    rule.line.fill.background()


def add_card(slide, x, y, w, h, heading, body, accent=TEAL, body_size=13):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                  Inches(w), Inches(h))
    set_fill(card, WHITE)
    set_line(card, LINE, 0.8)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                  Inches(0.06), Inches(h))
    set_fill(band, accent)
    band.line.fill.background()
    add_text(slide, heading, x + 0.20, y + 0.18, w - 0.35, 0.30, 13, INK, True)
    add_text(slide, body, x + 0.20, y + 0.57, w - 0.35, h - 0.70, body_size, MUTED)
    return card


def add_metric(slide, x, y, w, value, label, accent=TEAL):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                  Inches(w), Inches(1.12))
    set_fill(card, WHITE)
    set_line(card, LINE, 0.8)
    add_text(slide, value, x + 0.18, y + 0.17, w - 0.3, 0.40, 25, accent, True)
    add_text(slide, label, x + 0.18, y + 0.68, w - 0.3, 0.24, 10, MUTED, True)


def new_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(PAPER)
    return slide


def build_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Cover
    slide = new_slide(prs)
    block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                                   Inches(13.333), Inches(7.5))
    set_fill(block, INK)
    block.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(0.95),
                                    Inches(0.13), Inches(4.28))
    set_fill(accent, GOLD)
    accent.line.fill.background()
    add_text(slide, "AI ENGINEER INTERNSHIP CASE STUDY", 1.08, 1.02, 6.5, 0.3, 11, "B9E1D6", True)
    add_text(slide, "Document Intelligence\nPlatform", 1.05, 1.62, 8.8, 1.45, 38, WHITE, True)
    add_text(slide, "End-to-end extraction, validation and reconciliation for financial documents",
             1.08, 3.42, 8.5, 0.55, 17, "D8EAE5")
    add_text(slide, "Live application", 1.08, 5.18, 1.6, 0.22, 10, "A9C4BE", True)
    add_text(slide, "document-intelligence-platform-kace.onrender.com", 1.08, 5.47, 6.5, 0.28, 13, WHITE)
    add_text(slide, "FastAPI  |  Gemini vision  |  Evidence-backed JSON", 1.08, 6.56, 7.5, 0.25, 11, "A9C4BE")
    add_text(slide, "01", 11.85, 6.56, 0.7, 0.3, 12, "B9E1D6", True, align=PP_ALIGN.RIGHT)

    # 2. Objective and scope
    slide = new_slide(prs)
    add_title(slide, "Objective and scope", "Case-study requirements", 2)
    add_text(slide, "A reviewable pipeline for turning mixed financial documents into traceable, validated data.",
             0.68, 1.72, 11.3, 0.36, 16, MUTED)
    add_card(slide, 0.70, 2.35, 2.85, 1.65, "4 document types",
             "Invoice, balance sheet, profit and loss, and cash-flow statement.")
    add_card(slide, 3.78, 2.35, 2.85, 1.65, "Safe inputs",
             "PDF, JPG/JPEG and PNG; validated from bytes with a three-page limit.", GOLD)
    add_card(slide, 6.86, 2.35, 2.85, 1.65, "Traceable output",
             "Fields, table values, source snippets and page evidence.")
    add_card(slide, 9.94, 2.35, 2.85, 1.65, "Arithmetic checks",
             "Document-type formulas with tolerance and explicit statuses.", GOLD)
    add_text(slide, "Design principles", 0.70, 4.62, 2.3, 0.28, 14, TEAL, True)
    add_bullets(slide, [
        "Do not hallucinate: absent or unreadable values stay null.",
        "Reject invalid files before OCR or paid AI work.",
        "Keep asynchronous processing visible to the user.",
        "Expose the same structured result through the dashboard and API.",
    ], 0.72, 5.02, 11.2, 1.55, 15)

    # 3. Architecture
    slide = new_slide(prs)
    add_title(slide, "System architecture", "Modular processing pipeline", 3)
    image = DOCS / "architecture.png"
    if image.exists():
        slide.shapes.add_picture(str(image), Inches(5.45), Inches(1.72), width=Inches(7.25))
    add_text(slide, "Request flow", 0.70, 1.85, 2.0, 0.28, 14, TEAL, True)
    add_rich_lines(slide, [
        "Browser dashboard",
        "        -> FastAPI REST API",
        "        -> validation and OCR",
        "        -> Gemini vision extraction",
        "        -> financial validation",
        "        -> repository and database",
        "        -> dashboard / API response",
    ], 0.72, 2.30, 4.25, 2.65, 15, INK)
    add_text(slide, "Service boundaries", 0.70, 5.20, 2.3, 0.28, 14, TEAL, True)
    add_bullets(slide, [
        "document_validation_service.py",
        "ocr_service.py and extraction_service.py",
        "financial_validation_service.py",
        "document_repository.py",
    ], 0.72, 5.58, 4.2, 1.15, 12, MUTED)

    # 4. Validation and OCR
    slide = new_slide(prs)
    add_title(slide, "Validation, OCR and input safety", "Before the AI call", 4)
    add_card(slide, 0.70, 1.85, 3.75, 1.55, "1. Validate bytes",
             "Detect file signatures, MIME/type, integrity, readability, page count and size.")
    add_card(slide, 4.78, 1.85, 3.75, 1.55, "2. Extract text",
             "Use PDF text when available; rasterize scanned pages with PyMuPDF and Tesseract.", GOLD)
    add_card(slide, 8.86, 1.85, 3.75, 1.55, "3. Preserve evidence",
             "Normalize page-tagged text and images so every field can point back to a page.")
    add_text(slide, "Rejected early", 0.72, 4.08, 2.0, 0.28, 14, RED, True)
    add_bullets(slide, [
        "Unsupported extensions or content types",
        "Empty, corrupt or unreadable files",
        "Documents above 15 MB or more than 3 pages",
        "Invalid PDF/image content before OCR begins",
    ], 0.72, 4.48, 5.45, 1.52, 15)
    add_text(slide, "Supported upload contract", 7.18, 4.08, 3.0, 0.28, 14, TEAL, True)
    add_rich_lines(slide, [
        "PDF  |  JPG / JPEG  |  PNG",
        "Maximum 3 pages  |  Maximum 15 MB",
        "Processing states: PROCESSING -> PASS / FAILED",
        "Page number and source snippet per extracted field",
    ], 7.18, 4.48, 5.15, 1.6, 15, INK)

    # 5. AI extraction
    slide = new_slide(prs)
    add_title(slide, "Evidence-backed AI extraction", "Current Gemini implementation", 5)
    add_text(slide, "The model is used for reading and structuring the document, not for inventing accounting data.",
             0.70, 1.72, 11.4, 0.36, 16, MUTED)
    add_card(slide, 0.70, 2.35, 3.75, 2.15, "Structured contract",
             "google-genai SDK with a stable extracted_data envelope, line items, evidence and confidence.")
    add_card(slide, 4.78, 2.35, 3.75, 2.15, "Vision + text context",
             "Normalized page images plus page-tagged OCR/text are supplied for layout-aware extraction.", GOLD)
    add_card(slide, 8.86, 2.35, 3.75, 2.15, "Resilient responses",
             "Handles current response variants, empty output, JSON fences, retry and model fallback.")
    add_text(slide, "Guardrails", 0.70, 5.02, 1.5, 0.28, 14, TEAL, True)
    add_bullets(slide, [
        "Missing or unreadable values return null.",
        "Evidence includes source_text and page_number.",
        "Comparative periods use explicit field__period keys.",
        "Runtime key is configured through GEMINI_API_KEY and never committed.",
    ], 0.72, 5.40, 11.25, 1.15, 15)

    # 6. Financial validation
    slide = new_slide(prs)
    add_title(slide, "Financial validation engine", "Transparent arithmetic", 6)
    add_text(slide, "Each check records its operands, formula, calculated value, reported value, variance and status.",
             0.70, 1.72, 11.6, 0.35, 16, MUTED)
    checks = [
        ("Invoice", "subtotal + tax - discount = total", TEAL),
        ("Balance sheet", "liabilities + equity = assets", GOLD),
        ("Profit and loss", "revenue - cost of sales = gross profit", TEAL),
        ("Cash flow", "operating + investing + financing + FX = net change", GOLD),
    ]
    y = 2.35
    for label, formula, accent_color in checks:
        add_card(slide, 0.70, y, 11.9, 0.72, label, formula, accent_color, 14)
        y += 0.88
    add_text(slide, "Status policy", 0.70, 6.12, 1.75, 0.24, 14, TEAL, True)
    add_text(slide, "PASS within rounding tolerance  |  FAIL when operands disagree  |  NOT_APPLICABLE when inputs are missing",
             2.08, 6.12, 10.1, 0.27, 13, INK)

    # 7. Product experience and API
    slide = new_slide(prs)
    add_title(slide, "Product experience and API", "Reviewable by a human", 7)
    add_card(slide, 0.70, 1.85, 3.75, 2.05, "Ledger dashboard",
             "Responsive upload workflow, search, filters, statistics and refresh control.")
    add_card(slide, 4.78, 1.85, 3.75, 2.05, "Document workspace",
             "Centered detail view with Fields, Validation and Raw JSON tabs; missing and low-confidence values are marked.", GOLD)
    add_card(slide, 8.86, 1.85, 3.75, 2.05, "Async visibility",
             "POST returns a processing record; the browser polls until PASS or FAILED and shows the exact failure.")
    add_text(slide, "Mandatory endpoints", 0.70, 4.45, 2.6, 0.28, 14, TEAL, True)
    add_rich_lines(slide, [
        "POST  /api/v1/documents/process",
        "GET   /api/v1/documents/{document_name}",
        "GET   /api/v1/documents",
        "GET   /api/v1/health",
        "GET   /docs  and  /redoc",
    ], 0.72, 4.85, 6.3, 1.55, 15, INK)
    add_text(slide, "One response model", 8.86, 4.45, 2.8, 0.28, 14, TEAL, True)
    add_text(slide, "The API and UI share the same persisted document, extracted fields, line items, validation checks and raw JSON.",
             8.88, 4.85, 3.62, 1.28, 15, INK)

    # 8. Verification
    slide = new_slide(prs)
    add_title(slide, "Verification and deployment", "Evidence of completion", 8)
    add_metric(slide, 0.70, 1.90, 2.7, "25", "AUTOMATED TESTS PASS")
    add_metric(slide, 3.70, 1.90, 2.7, "43", "LIVE FIELDS EXTRACTED", GOLD)
    add_metric(slide, 6.70, 1.90, 2.7, "6", "LIVE LINE ITEMS", TEAL)
    add_metric(slide, 9.70, 1.90, 2.9, "PASS", "FRESH LIVE INVOICE TEST", GOLD)
    add_text(slide, "Deployment checks", 0.70, 3.55, 2.2, 0.28, 14, TEAL, True)
    add_bullets(slide, [
        "Public Render frontend is live.",
        "/api/v1/health reports service status and Gemini configuration.",
        "Swagger/OpenAPI is available at /docs.",
        "Fresh live invoice verification returned file validation PASS and processing PASS.",
        "Sample JSON fixtures cover all four document types and failure paths.",
    ], 0.72, 3.95, 7.25, 2.05, 15)
    add_text(slide, "Repository", 9.20, 3.55, 1.5, 0.28, 14, TEAL, True)
    add_text(slide, "github.com/SakthiPriyaR/\ndocument-intelligence-platform", 9.20, 3.95, 3.0, 0.72, 15, INK, True)
    add_text(slide, "No secrets committed; runtime key is supplied by the deployment environment.",
             9.20, 5.08, 3.0, 0.72, 13, MUTED)

    # 9. Deployment and submission links
    slide = new_slide(prs)
    add_title(slide, "Deployment and submission links", "Ready to submit", 9)
    add_text(slide, "Use these live links in the internship submission form and during your presentation.",
             0.70, 1.72, 11.6, 0.35, 16, MUTED)
    add_card(slide, 0.70, 2.28, 5.70, 1.25, "Deployed Backend API (health)",
             "https://document-intelligence-platform-kace.onrender.com/api/v1/health", TEAL, 10)
    add_link(slide, "Open working API endpoint", "https://document-intelligence-platform-kace.onrender.com/api/v1/health",
             0.95, 3.06, 2.8, 0.22, 11)
    add_card(slide, 6.72, 2.28, 5.88, 1.25, "Deployed Frontend URL",
             "https://document-intelligence-platform-kace.onrender.com/", GOLD, 11)
    add_link(slide, "Open live frontend", "https://document-intelligence-platform-kace.onrender.com/",
             6.97, 3.06, 2.2, 0.22, 11)
    add_card(slide, 0.70, 3.78, 5.70, 1.25, "Deployment platform used",
             "Render - Docker web service on the Free plan", TEAL, 13)
    add_card(slide, 6.72, 3.78, 5.88, 1.25, "Public GitHub repository",
             "github.com/SakthiPriyaR/document-intelligence-platform", GOLD, 12)
    add_link(slide, "Open public repository", "https://github.com/SakthiPriyaR/document-intelligence-platform",
             6.97, 4.56, 2.8, 0.22, 11)
    add_text(slide, "PPT sharing", 0.70, 5.63, 1.5, 0.26, 14, TEAL, True)
    add_link(slide, "Download the version-controlled PPTX from GitHub",
             "https://github.com/SakthiPriyaR/document-intelligence-platform/blob/main/docs/solution_presentation.pptx",
             0.72, 6.02, 4.6, 0.26, 12)
    add_text(slide, "Google Drive submission: upload this PPTX to your Drive, set General access to 'Anyone with the link - Viewer', then paste that share URL into the form.",
             5.58, 5.92, 6.65, 0.55, 13, INK)

    # 10. Limitations
    slide = new_slide(prs)
    add_title(slide, "Limitations and production plan", "What changes beyond evaluation", 10)
    add_text(slide, "The evaluation deployment is intentionally small; the seams for production hardening are explicit.",
             0.70, 1.72, 11.3, 0.35, 16, MUTED)
    add_card(slide, 0.70, 2.35, 5.65, 2.65, "Current limitations",
             "Render Free local SQLite storage is ephemeral.\n\nNo authentication or rate limiting.\n\nLarge dense documents may need chunked extraction.\n\nModel confidence is not independently calibrated.", RED, 14)
    add_card(slide, 6.70, 2.35, 5.9, 2.65, "Production changes",
             "Managed Postgres and object storage.\n\nDurable worker queue and observability.\n\nAuthentication, tenant isolation and quotas.\n\nGolden-document evaluations and reconciliation expansion.", TEAL, 14)
    add_text(slide, "Important distinction", 0.70, 5.62, 2.6, 0.28, 14, GOLD, True)
    add_text(slide, "The current live extraction flow works; durable history requires a persistent database plan.",
             0.72, 6.00, 11.3, 0.35, 16, INK)

    # 11. AI usage and deliverables
    slide = new_slide(prs)
    add_title(slide, "AI usage and submission deliverables", "Transparent engineering", 11)
    add_card(slide, 0.70, 1.85, 5.65, 3.25, "AI usage declaration",
             "An AI coding assistant supported the initial scaffold, service boundaries, validation engine, frontend, tests, documentation and this presentation.\n\nThe source was reviewed and locally test-verified. Gemini is used at runtime only for evidence-backed field extraction after explicit configuration.", TEAL, 14)
    add_card(slide, 6.70, 1.85, 5.9, 3.25, "Included in the repository",
             "Public GitHub source\nLive frontend, API and Swagger\nREADME with setup and limitations\nArchitecture source plus PNG/PDF exports\nThis solution presentation in PPTX/PDF/Markdown\nSample fixtures and automated tests", GOLD, 14)
    add_text(slide, "Thank you", 0.70, 5.82, 2.4, 0.48, 25, TEAL, True)
    add_text(slide, "Live app: document-intelligence-platform-kace.onrender.com", 0.72, 6.35, 7.6, 0.27, 14, INK)

    # 12. Demo walkthrough
    slide = new_slide(prs)
    add_title(slide, "How to present the live demo", "Six-minute walkthrough", 12)
    steps = [
        ("01", "Open the live Ledger dashboard", "State the problem: financial documents are unstructured and hard to review."),
        ("02", "Choose a document type and upload", "Point out PDF/JPG/PNG support, page limit and validation before AI."),
        ("03", "Explain the processing path", "Validation -> OCR -> Gemini extraction -> financial reconciliation -> persisted result."),
        ("04", "Open the processed document", "Show fields, evidence, line items, validation checks and raw JSON."),
        ("05", "Close with assurance", "Mention live deployment, API docs, 25 tests, known limits and the production roadmap."),
    ]
    y = 1.72
    for number, heading, talk_track in steps:
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72), Inches(y), Inches(0.48), Inches(0.48))
        set_fill(badge, TEAL)
        badge.line.fill.background()
        add_text(slide, number, 0.72, y + 0.12, 0.48, 0.18, 9, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(slide, heading, 1.40, y + 0.03, 4.15, 0.25, 14, INK, True)
        add_text(slide, talk_track, 5.60, y + 0.03, 6.65, 0.38, 13, MUTED)
        y += 0.88
    add_text(slide, "Tip: do the upload before your interview starts so the completed record is already available to open.",
             0.72, 6.35, 10.8, 0.28, 13, GOLD, True)

    prs.save(path)


def pdf_text(c, text: str, x: float, y: float, width: int, size=14,
             color=colors.HexColor("#102C32"), leading=19, bold=False):
    c.setFillColor(color)
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    for line in text.splitlines():
        pieces = wrap(line, width=width) or [""]
        for piece in pieces:
            c.drawString(x, y, piece)
            y -= leading
    return y


def build_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=landscape(letter))
    width, height = landscape(letter)
    pages = [
        ("Document Intelligence Platform", "AI Engineer Internship Case Study", [
            "End-to-end extraction, validation and reconciliation for financial documents.",
            "Live application: document-intelligence-platform-kace.onrender.com",
            "FastAPI | Gemini vision | Evidence-backed JSON",
        ]),
        ("Objective and scope", "Case-study requirements", [
            "Four document types: invoice, balance sheet, profit and loss, cash flow statement.",
            "PDF/JPG/PNG uploads, maximum three pages and 15 MB, validated before OCR or AI.",
            "Meaningful fields and tables with source snippets and page evidence.",
            "Document-type financial formulas with PASS, FAIL and NOT_APPLICABLE.",
        ]),
        ("System architecture", "Modular processing pipeline", [
            "Browser dashboard -> FastAPI -> validation -> OCR -> Gemini vision -> financial validation -> repository -> database.",
            "See architecture.png and architecture.pdf for the complete flow and service boundaries.",
        ]),
        ("Validation, OCR and input safety", "Before the AI call", [
            "File signatures, type, integrity, readability, page count and size are checked from bytes.",
            "Native PDF text is preferred; scanned pages use PyMuPDF rasterization and Tesseract fallback.",
            "Page-tagged text and normalized images preserve provenance for every field.",
        ]),
        ("Evidence-backed AI extraction", "Current Gemini implementation", [
            "google-genai SDK, strict JSON envelope, line items, evidence and confidence.",
            "Missing or unreadable values stay null; no unsupported inference.",
            "Response variants, empty output, JSON fences, retry and model fallback are handled.",
        ]),
        ("Financial validation engine", "Transparent arithmetic", [
            "Invoice: subtotal + tax - discount = total.",
            "Balance sheet: liabilities + equity = assets.",
            "Profit and loss: revenue - cost of sales = gross profit.",
            "Cash flow: operating + investing + financing + FX = net change.",
            "Rounding tolerance is explicit; missing operands produce NOT_APPLICABLE.",
        ]),
        ("Product experience and API", "Reviewable by a human", [
            "Responsive dashboard, upload workflow, search, filters, statistics and refresh.",
            "Centered document workspace with Fields, Validation and Raw JSON tabs.",
            "POST /api/v1/documents/process; GET /api/v1/documents; GET /api/v1/documents/{name}; GET /api/v1/health; /docs.",
        ]),
        ("Verification and deployment", "Evidence of completion", [
            "25 automated tests pass without a network call or API key.",
            "Fresh live invoice verification returned file validation PASS and processing PASS.",
            "The live test returned 43 extracted field entries and 6 line items.",
            "Public Render deployment, health endpoint, Swagger and four-type sample fixtures are included.",
        ]),
        ("Deployment and submission links", "Ready to submit", [
            "Backend API health endpoint: https://document-intelligence-platform-kace.onrender.com/api/v1/health (API prefix: /api/v1).",
            "Frontend URL: https://document-intelligence-platform-kace.onrender.com/",
            "Deployment platform: Render Docker web service.",
            "Public repository: https://github.com/SakthiPriyaR/document-intelligence-platform",
            "For a Google Drive PPT link, upload solution_presentation.pptx and set access to Anyone with the link - Viewer.",
        ]),
        ("Limitations and production plan", "What changes beyond evaluation", [
            "Render Free local SQLite storage is ephemeral; managed Postgres is required for durable history.",
            "Production adds authentication, tenant isolation, object storage, durable workers, observability and evaluations.",
        ]),
        ("AI usage and submission deliverables", "Transparent engineering", [
            "AI coding assistance was reviewed and test-verified; Gemini is used only for configured, evidence-backed extraction.",
            "Repository includes source, live links, README, tests, samples, architecture source/PNG/PDF and this presentation.",
        ]),
        ("How to present the live demo", "Six-minute walkthrough", [
            "Introduce the problem: turn unstructured financial files into traceable ledger records.",
            "Choose a document type, upload a supported document and explain validation before AI.",
            "Describe the path: OCR, Gemini extraction, arithmetic checks and persistence.",
            "Open the processed result to show fields, evidence, line items, validation and raw JSON.",
            "Close with deployment links, test evidence, limitations and production roadmap.",
        ]),
    ]
    architecture = DOCS / "architecture.png"
    for index, (title, kicker, bullets) in enumerate(pages, start=1):
        c.setFillColor(colors.HexColor("#F6F4EE"))
        c.rect(0, 0, width, height, fill=1, stroke=0)
        c.setFillColor(colors.HexColor("#0F766E"))
        c.rect(42, height - 58, 5, 30, fill=1, stroke=0)
        pdf_text(c, kicker.upper(), 58, height - 42, 100, 9, colors.HexColor("#0F766E"), 11, True)
        pdf_text(c, title, 58, height - 82, 70, 25, colors.HexColor("#102C32"), 29, True)
        c.setStrokeColor(colors.HexColor("#D6E1DE"))
        c.line(58, height - 105, width - 58, height - 105)
        if index == 3 and architecture.exists():
            c.drawImage(str(architecture), 425, 140, width=350, height=225, preserveAspectRatio=True, anchor="c")
            pdf_text(c, bullets[0], 58, height - 150, 48, 14, colors.HexColor("#102C32"), 20)
            pdf_text(c, bullets[1], 58, height - 275, 48, 14, colors.HexColor("#66767A"), 20)
        else:
            y = height - 155
            for bullet in bullets:
                c.setFillColor(colors.HexColor("#C48A38"))
                c.circle(72, y + 4, 3, fill=1, stroke=0)
                y = pdf_text(c, bullet, 90, y, 82, 16, colors.HexColor("#102C32"), 24)
                y -= 12
        c.setFillColor(colors.HexColor("#66767A"))
        c.setFont("Helvetica-Bold", 8)
        c.drawRightString(width - 58, 30, f"{index:02d}  |  DOCUMENT INTELLIGENCE PLATFORM")
        c.showPage()
    c.save()


def build() -> None:
    build_pptx(DOCS / "solution_presentation.pptx")
    build_pdf(DOCS / "solution_presentation.pdf")


if __name__ == "__main__":
    build()
